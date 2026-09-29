"""Erzeugt vorlagen/skript-vorlage.docx, die Word-Referenzvorlage für die Skripte.

Aufruf im Projektordner:  python vorlagen/erzeuge_word_vorlage.py
Benötigt: python-docx und Quarto (für die Pandoc-Standardvorlage).

Die Vorlage kann auch direkt in Word angepasst werden (Kopf-/Fußzeile,
Formatvorlagen „Heading 1–3“, „Title“, „Body Text“). Pandoc übernimmt Seiten-
layout, Kopf- und Fußzeilen und Formatvorlagen, nicht den Beispielinhalt.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ZIEL = HERE / "skript-vorlage.docx"
LOGO = HERE / "hsm-wirtschaft-logo.png"
BLAU = RGBColor(0x2F, 0x54, 0x96)
FUSSZEILE = "DIRK LOOMANS · VORLESUNGSSKRIPT GRUNDLAGEN DER RESILIENZ · HOCHSCHULE MAINZ"


def standardvorlage():
    daten = subprocess.run(
        ["quarto", "pandoc", "-o", str(ZIEL), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    return daten


def feld(run, anweisung):
    for typ, text in (("begin", None), (None, anweisung), ("end", None)):
        if typ:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), typ)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def obere_linie(absatz, farbe="4472C4", staerke=18):
    ppr = absatz._p.get_or_add_pPr()
    rand = OxmlElement("w:pBdr")
    oben = OxmlElement("w:top")
    for k, v in (("w:val", "single"), ("w:sz", str(staerke)), ("w:space", "4"), ("w:color", farbe)):
        oben.set(qn(k), v)
    rand.append(oben)
    ppr.append(rand)


def main():
    standardvorlage()
    doc = Document(ZIEL)

    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin, sec.bottom_margin = Cm(3.2), Cm(2.6)
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)
        sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

        kopf = sec.header.paragraphs[0]
        kopf.text = ""
        kopf.alignment = 2  # rechts
        kopf.add_run().add_picture(str(LOGO), height=Cm(1.25))

        fuss = sec.footer.paragraphs[0]
        fuss.text = ""
        obere_linie(fuss)
        breite = sec.page_width - sec.left_margin - sec.right_margin
        fuss.paragraph_format.tab_stops.add_tab_stop(breite, WD_TAB_ALIGNMENT.RIGHT)
        r = fuss.add_run(FUSSZEILE + "\t")
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        seite = fuss.add_run()
        seite.font.size = Pt(8)
        feld(seite, "PAGE")

    stile = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in stile]:
            st = stile[name]
            st.font.name = "Calibri"
            st.font.size = Pt(10.5)
    for name, groesse in (("Title", 20), ("Subtitle", 13), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        st = stile[name]
        st.font.name = "Calibri"
        st.font.size = Pt(groesse)
        st.font.color.rgb = BLAU
        st.font.bold = name != "Subtitle"
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts.attrib.pop(qn(attr), None)
    doc.save(ZIEL)
    print("Vorlage geschrieben:", ZIEL)


if __name__ == "__main__":
    main()
